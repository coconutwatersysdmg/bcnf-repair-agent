from __future__ import annotations

import csv
import math
import os
import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"E:\Homework\1研究生课程\AAA小论文\A论文编写\研究方案与实验项目_20261017")
REPO = ROOT / "repair_agent_experiment" / ".worktrees" / "experiment-pilot-20261021"
REFERENCE = ROOT / "BCNF条件下的数据库子集修复检查与大模型辅助修复_理论保留版_20261021.docx"
OUTPUT_DIR = ROOT / "论文成稿"
OUTPUT = OUTPUT_DIR / "BCNF条件下的子集修复检查与大模型反例反馈方法_实验增强稿_20261025.docx"
ASSET_DIR = ROOT / ".codex_docx_build_20261025" / "figures"

MAIN_CSV = REPO / "reports" / "tables" / "main_results.csv"
EFFECT_CSV = REPO / "reports" / "tables" / "paired_effects.csv"
COST_CSV = REPO / "reports" / "tables" / "cost_latency.csv"
BENCH_CSV = REPO / "results" / "benchmarks" / "bcnf_backend_20261024" / "bcnf_backend_summary.csv"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


MAIN = read_csv(MAIN_CSV)
EFFECTS = read_csv(EFFECT_CSV)
COSTS = read_csv(COST_CSV)
BENCH = read_csv(BENCH_CSV)


DATASET_LABEL = {
    "synthetic_confirmation": "合成确认集",
    "semireal_bike": "Bike Sharing",
    "semireal_wdbc": "WDBC",
}
ARM_LABEL = {
    "D0": "D0 确定性贪心",
    "D1": "D1 最大基数参考",
    "B0": "B0 单轮生成",
    "B1": "B1 布尔反馈",
    "B2": "B2 完整反例",
    "B3_192": "B3 覆盖优先 192 B",
    "B3_truncate_192": "截断 192 B",
    "B3_random_192": "随机 192 B",
    "B3_384": "B3 覆盖优先 384 B",
    "B3_truncate_384": "截断 384 B",
    "B3_random_384": "随机 384 B",
}


def find_main(dataset: str, arm: str) -> dict[str, str]:
    return next(r for r in MAIN if r["dataset"] == dataset and r["arm"] == arm)


def find_cost(dataset: str, arm: str) -> dict[str, str]:
    return next(r for r in COSTS if r["dataset"] == dataset and r["arm"] == arm)


def find_effect(dataset: str, comparison: str, metric: str = "repair_rate") -> dict[str, str]:
    return next(
        r for r in EFFECTS
        if r["dataset"] == dataset and r["comparison"] == comparison and r["metric"] == metric
    )


def pct(value: str | float, digits: int = 1) -> str:
    return f"{float(value) * 100:.{digits}f}"


def ms(ns: str | float) -> float:
    return float(ns) / 1_000_000.0


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    paths = [
        Path(r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf" if bold else r"C:\Windows\Fonts\simsun.ttc"),
        Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
    ]
    for p in paths:
        if p.exists():
            return ImageFont.truetype(str(p), size=size)
    return ImageFont.load_default()


NAVY = "#17365D"
BLUE = "#2F75B5"
TEAL = "#2A9D8F"
ORANGE = "#E07A5F"
GOLD = "#D4A72C"
SLATE = "#65758B"
LIGHT = "#E9EFF6"
GRID = "#D8DEE8"
TEXT = "#263238"


def base_canvas(title: str, width: int = 2200, height: int = 1280):
    im = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(im)
    d.text((width // 2, 55), title, font=font(46, True), fill=NAVY, anchor="ma")
    return im, d


def draw_axes(d: ImageDraw.ImageDraw, box: tuple[int, int, int, int], y_max: float, y_label: str, ticks: int = 5):
    x0, y0, x1, y1 = box
    d.line((x0, y0, x0, y1), fill=TEXT, width=3)
    d.line((x0, y1, x1, y1), fill=TEXT, width=3)
    for i in range(ticks + 1):
        v = y_max * i / ticks
        y = y1 - (y1 - y0) * i / ticks
        d.line((x0, y, x1, y), fill=GRID, width=2)
        d.text((x0 - 20, y), f"{v:.0f}", font=font(25), fill=TEXT, anchor="rm")
    d.text((45, (y0 + y1) // 2), y_label, font=font(29, True), fill=TEXT, anchor="mm")


def make_repair_chart(path: Path):
    im, d = base_canvas("三组数据上的形式修复率")
    box = (170, 210, 2110, 1020)
    draw_axes(d, box, 100, "修复率 / %")
    datasets = list(DATASET_LABEL)
    arms = ["D0", "B0", "B1", "B2", "B3_192", "B3_384"]
    colors = [NAVY, SLATE, ORANGE, TEAL, GOLD, BLUE]
    group_w = (box[2] - box[0]) / len(datasets)
    bar_w = 68
    gap = 13
    for gi, ds in enumerate(datasets):
        cx = box[0] + group_w * (gi + 0.5)
        total = len(arms) * bar_w + (len(arms) - 1) * gap
        start = cx - total / 2
        for ai, arm in enumerate(arms):
            value = float(find_main(ds, arm)["repair_rate"]) * 100
            x0 = start + ai * (bar_w + gap)
            y = box[3] - (box[3] - box[1]) * value / 100
            d.rounded_rectangle((x0, y, x0 + bar_w, box[3]), radius=8, fill=colors[ai])
            d.text((x0 + bar_w / 2, y - 10), f"{value:.0f}", font=font(20, True), fill=TEXT, anchor="mb")
        d.text((cx, box[3] + 40), DATASET_LABEL[ds], font=font(29, True), fill=TEXT, anchor="ma")
    lx, ly = 210, 1120
    for i, arm in enumerate(arms):
        x = lx + (i % 3) * 610
        y = ly + (i // 3) * 62
        d.rounded_rectangle((x, y, x + 45, y + 30), radius=5, fill=colors[i])
        d.text((x + 60, y + 15), ARM_LABEL[arm], font=font(25), fill=TEXT, anchor="lm")
    im.save(path, dpi=(220, 220))


def make_effect_chart(path: Path):
    rows = []
    for ds in DATASET_LABEL:
        rows.append((f"{DATASET_LABEL[ds]}  B2−B1", find_effect(ds, "B2-B1"), TEAL))
    for ds in DATASET_LABEL:
        rows.append((f"{DATASET_LABEL[ds]}  B2: C1−C0", find_effect(ds, "B2:C1-C0"), BLUE))
    im, d = base_canvas("配对效应与 95% 置信区间", height=1000)
    x0, x1 = 720, 2070
    y0, y1 = 165, 790
    lo, hi = -12.0, 55.0
    def xp(v): return x0 + (v - lo) / (hi - lo) * (x1 - x0)
    for tick in [-10, 0, 10, 20, 30, 40, 50]:
        x = xp(tick)
        d.line((x, y0, x, y1), fill="#8795A8" if tick == 0 else GRID, width=5 if tick == 0 else 2)
        d.text((x, y1 + 35), f"{tick:+d}", font=font(25), fill=TEXT, anchor="ma")
    for i, (label, r, color) in enumerate(rows):
        y = y0 + 45 + i * 96
        eff = float(r["absolute_difference"]) * 100
        low = float(r["ci95_low"]) * 100
        high = float(r["ci95_high"]) * 100
        d.text((675, y), label, font=font(27), fill=TEXT, anchor="rm")
        d.line((xp(low), y, xp(high), y), fill=color, width=9)
        d.ellipse((xp(eff)-13, y-13, xp(eff)+13, y+13), fill=color, outline="white", width=3)
    d.text(((x0+x1)//2, 900), "修复率差异 / 百分点", font=font(30, True), fill=TEXT, anchor="ma")
    im.save(path, dpi=(220, 220))


def make_budget_chart(path: Path):
    im, d = base_canvas("覆盖优先策略相对预算对照的修复率差异")
    box = (210, 230, 2100, 990)
    x0, x1, y0, y1 = box[0], box[2], box[1], box[3]
    lo, hi = -4.0, 4.0
    def yp(v): return y1 - (v - lo) / (hi - lo) * (y1 - y0)
    for tick in [-4, -2, 0, 2, 4]:
        y = yp(tick)
        d.line((x0, y, x1, y), fill="#8795A8" if tick == 0 else GRID, width=5 if tick == 0 else 2)
        d.text((x0-18, y), f"{tick:+d}", font=font(25), fill=TEXT, anchor="rm")
    datasets = list(DATASET_LABEL)
    comparisons = [
        ("B3_192-B3_truncate_192", "192 B−截断", BLUE),
        ("B3_192-B3_random_192", "192 B−随机", TEAL),
        ("B3_384-B3_truncate_384", "384 B−截断", ORANGE),
        ("B3_384-B3_random_384", "384 B−随机", GOLD),
    ]
    group_w = (x1-x0)/3
    bw = 82
    for gi, ds in enumerate(datasets):
        cx = x0 + group_w*(gi+0.5)
        start = cx - 1.5*(bw+22)
        for ai,(comp,_,color) in enumerate(comparisons):
            val=float(find_effect(ds,comp)["absolute_difference"])*100
            x=start+ai*(bw+22)
            y=yp(val)
            top=min(y,yp(0)); bottom=max(y,yp(0))
            if abs(bottom-top)<5: bottom=top+5
            d.rounded_rectangle((x,top,x+bw,bottom),radius=6,fill=color)
            d.text((x+bw/2, y-10 if val>=0 else y+10),f"{val:+.1f}",font=font(19,True),fill=TEXT,anchor="mb" if val>=0 else "ma")
        d.text((cx,y1+45),DATASET_LABEL[ds],font=font(29,True),fill=TEXT,anchor="ma")
    for i,(_,label,color) in enumerate(comparisons):
        x=280+(i%2)*900; y=1100+(i//2)*60
        d.rounded_rectangle((x,y,x+45,y+30),radius=5,fill=color)
        d.text((x+60,y+15),label,font=font(25),fill=TEXT,anchor="lm")
    d.text((55,(y0+y1)//2),"修复率差异 / 百分点",font=font(29,True),fill=TEXT,anchor="mm")
    im.save(path,dpi=(220,220))


def make_scaling_chart(path: Path):
    selected = [r for r in BENCH if r["family"] == "composite" and r["conflict_ratio"] == "0.05"]
    counts = sorted({int(r["row_count"]) for r in selected})
    algorithms = [
        ("bcnf_index", "BCNF共享键索引", NAVY),
        ("fd_hash", "通用FD哈希", BLUE),
        ("incremental", "增量回放", TEAL),
        ("repair_agent_generic", "通用检查后端", ORANGE),
    ]
    im, d = base_canvas("BCNF检查后端的规模扩展")
    box=(190,220,2100,1000)
    maxv=max(ms(r["median_decision_ns"]) for r in selected if r["algorithm"] in {a[0] for a in algorithms})
    ymax=math.ceil(maxv/50)*50
    draw_axes(d,box,ymax,"判定时间 / ms")
    xgap=(box[2]-box[0])/(len(counts)-1)
    for i,n in enumerate(counts):
        x=box[0]+i*xgap
        d.text((x,box[3]+42),f"{n//1000}k" if n>=1000 else str(n),font=font(28,True),fill=TEXT,anchor="ma")
    for alg,label,color in algorithms:
        pts=[]
        for i,n in enumerate(counts):
            r=next((x for x in selected if x["algorithm"]==alg and int(x["row_count"])==n),None)
            if not r: continue
            v=ms(r["median_decision_ns"])
            x=box[0]+i*xgap
            y=box[3]-(box[3]-box[1])*v/ymax
            pts.append((x,y))
        d.line(pts,fill=color,width=8,joint="curve")
        for x,y in pts: d.ellipse((x-12,y-12,x+12,y+12),fill=color,outline="white",width=3)
    for i,(_,label,color) in enumerate(algorithms):
        x=270+(i%2)*900; y=1090+(i//2)*58
        d.line((x,y+15,x+50,y+15),fill=color,width=8)
        d.ellipse((x+18,y+3,x+42,y+27),fill=color)
        d.text((x+70,y+15),label,font=font(25),fill=TEXT,anchor="lm")
    d.text(((box[0]+box[2])//2,box[3]+100),"记录数",font=font(30,True),fill=TEXT,anchor="ma")
    im.save(path,dpi=(220,220))


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    if size is not None: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color is not None: run.font.color.rgb = RGBColor.from_string(color)


def clear_paragraph(p):
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)


def set_paragraph_text(p, label: str, text: str, english: bool = False):
    clear_paragraph(p)
    r1=p.add_run(label)
    set_run_font(r1,10.5,True)
    r2=p.add_run(text)
    set_run_font(r2,10.5)
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent=Pt(0)
    p.paragraph_format.line_spacing=1.2
    p.paragraph_format.space_after=Pt(3)


def three_line_table(table, header_fill="D9E5F2"):
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=True
    tblPr=table._tbl.tblPr
    borders=tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders=OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ["top","bottom"]:
        el=OxmlElement(f"w:{edge}")
        el.set(qn("w:val"),"single"); el.set(qn("w:sz"),"12"); el.set(qn("w:color"),"17365D")
        borders.append(el)
    for edge in ["left","right","insideH","insideV"]:
        el=OxmlElement(f"w:{edge}")
        el.set(qn("w:val"),"nil")
        borders.append(el)
    for j,cell in enumerate(table.rows[0].cells):
        tcPr=cell._tc.get_or_add_tcPr()
        shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),header_fill); tcPr.append(shd)
        bottom=OxmlElement("w:tcBorders")
        b=OxmlElement("w:bottom"); b.set(qn("w:val"),"single"); b.set(qn("w:sz"),"8"); b.set(qn("w:color"),"17365D")
        bottom.append(b); tcPr.append(bottom)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
                for run in p.runs: set_run_font(run,8.2,row is table.rows[0])


def add_table(doc: Document, conclusion_p, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table=doc.add_table(rows=1,cols=len(headers))
    for i,h in enumerate(headers): table.rows[0].cells[i].text=h
    for data in rows:
        cells=table.add_row().cells
        for i,v in enumerate(data): cells[i].text=str(v)
    three_line_table(table)
    if widths:
        for row in table.rows:
            for i,w in enumerate(widths): row.cells[i].width=Cm(w)
    conclusion_p._p.addprevious(table._tbl)
    return table


def add_before(doc: Document, conclusion_p, text: str = "", style: str = "Normal", indent: bool = True):
    p=doc.add_paragraph(style=style)
    p.add_run(text)
    if style=="Normal":
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent=Pt(21 if indent else 0)
        p.paragraph_format.space_after=Pt(3)
        p.paragraph_format.line_spacing=1.2
        for run in p.runs: set_run_font(run,10.5)
    elif style.startswith("Heading"):
        p.paragraph_format.keep_with_next=True
    conclusion_p._p.addprevious(p._p)
    return p


def add_caption(doc: Document, conclusion_p, text: str):
    style="Figure Caption"
    p=add_before(doc,conclusion_p,text,style,False)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=not text.startswith("图")
    for run in p.runs: set_run_font(run,9.5)
    return p


def add_figure(doc: Document, conclusion_p, image_path: Path, caption: str, width_inches: float = 6.35):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(2)
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(image_path),width=Inches(width_inches))
    conclusion_p._p.addprevious(p._p)
    add_caption(doc,conclusion_p,caption)


def set_cell_margins(cell, top=70, start=80, bottom=70, end=80):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in("w:tcMar")
    if tcMar is None: tcMar=OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for m,v in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node=tcMar.find(qn(f"w:{m}"))
        if node is None: node=OxmlElement(f"w:{m}"); tcMar.append(node)
        node.set(qn("w:w"),str(v)); node.set(qn("w:type"),"dxa")


def build():
    OUTPUT_DIR.mkdir(parents=True,exist_ok=True)
    ASSET_DIR.mkdir(parents=True,exist_ok=True)
    make_repair_chart(ASSET_DIR/"fig_repair_rate.png")
    make_effect_chart(ASSET_DIR/"fig_paired_effects.png")
    make_budget_chart(ASSET_DIR/"fig_budget_effect.png")
    make_scaling_chart(ASSET_DIR/"fig_bcnf_scaling.png")

    shutil.copy2(REFERENCE,OUTPUT)
    doc=Document(OUTPUT)

    # Title and bilingual abstracts.
    doc.paragraphs[0].text="BCNF条件下的子集修复检查与大模型反例反馈方法"
    doc.paragraphs[1].text="Subset Repair Checking and LLM Counterexample Feedback under BCNF"
    for p in doc.paragraphs[:2]:
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    abstract_p=doc.paragraphs[2]
    authors=abstract_p.insert_paragraph_before("刘朔飞¹，齐子锋¹，刘国华¹*，李承烨¹，刘晓宇¹，张格¹，王森²，余意²")
    authors.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in authors.runs: set_run_font(r,10.5)
    aff=abstract_p.insert_paragraph_before("（1. 东华大学信息与智能科学学院，上海 201620；2. 上海宝信软件股份有限公司，上海 201999）")
    aff.alignment=WD_ALIGN_PARAGRAPH.CENTER
    aff.paragraph_format.space_after=Pt(6)
    for r in aff.runs: set_run_font(r,9.5)

    zh_abs=("子集修复检查要求候选子实例既满足函数依赖，又在包含关系下具有极大性。针对通用检查重复验证、"
            "大语言模型候选易残留冲突或过度删除的问题，本文研究BCNF关系模式下的子集修复检查及可验证反例反馈。"
            "在保留BCNF函数依赖集性质与判定理论的基础上，通过左部冗余属性约简和相同候选键归并构建共享索引，"
            "并把冲突元组对与可加回元组分别作为一致性和极大性反例反馈给模型。实验采用合成确认集、Bike Sharing与WDBC"
            "两套半真实数据，每套包含90个基础实例、C0/C1同底表任务和3次重复。完整反例反馈B2相对布尔反馈B1的形式修复率"
            "分别提高44.3、35.7和32.6个百分点，三组Holm校正后p值均为0.0032。192 B和384 B预算下，覆盖优先策略相对截断或随机"
            "均未显示稳定优势。BCNF专用后端在6000个候选状态上与通用FD和独立参考后端保持100%一致；在5万记录复合键场景中，"
            "中位判定时间为89.8 ms，较通用后端减少40.3%。结果表明，结构化反例主要改善形式合法性，尚不能替代事实证据或证明跨模型泛化。")
    set_paragraph_text(doc.paragraphs[4],"摘要：",zh_abs)
    set_paragraph_text(doc.paragraphs[5],"关键词：","关系数据库；BCNF；函数依赖；子集修复检查；反例反馈；大语言模型")
    en_abs=("Subset repair checking requires a candidate subinstance to satisfy functional dependencies and to be inclusion-maximal. "
            "This paper studies checking under BCNF and turns deterministic counterexamples into feedback for LLM-assisted repair. "
            "Redundant determinant attributes are reduced and identical candidate keys share indexes; conflicting tuple pairs and addable tuples explain consistency and maximality failures. "
            "The evaluation uses a synthetic confirmation set and two semi-real datasets derived from Bike Sharing and WDBC, each with 90 base instances, paired C0/C1 tasks, and three repetitions. "
            "Compared with binary feedback, full counterexamples improve formal repair rates by 44.3, 35.7, and 32.6 percentage points, with Holm-adjusted p=0.0032 in all three datasets. "
            "Coverage-oriented selection under 192-byte and 384-byte budgets does not show a stable advantage over truncation or random selection. "
            "The BCNF backend agrees with generic and independent reference checkers on all 6,000 states and reduces median decision time by 40.3% at 50,000 rows in the composite-key setting. "
            "The evidence supports improved formal validity, but not factual correctness or cross-model generalization.")
    set_paragraph_text(doc.paragraphs[6],"Abstract: ",en_abs,True)
    set_paragraph_text(doc.paragraphs[7],"Keywords: ","relational database; BCNF; functional dependency; subset repair checking; counterexample feedback; large language model",True)

    # Remove the former experimental plan and insert frozen results before the conclusion.
    p4=next(p for p in doc.paragraphs if p.text.strip()=="4 实验与分析")
    p5=next(p for p in doc.paragraphs if p.text.strip()=="5 结论")
    node=p4._p.getnext()
    while node is not None and node is not p5._p:
        nxt=node.getnext(); node.getparent().remove(node); node=nxt

    add_before(doc,p5,"4.1 实验设置", "Heading 2",False)
    add_before(doc,p5,"实验围绕三个问题展开：完整反例是否优于仅返回通过或失败的布尔反馈；有限字节预算下的见证选择是否优于简单截断和随机选择；BCNF专用后端能否在保持判定一致的同时降低检查开销。模型实验固定使用Qwen3.7-Flash-2026-07-15，提示版本、温度、最大轮数和输出协议在冻结矩阵内保持不变。D0为来源分数排序后的确定性极大子集，D1为小规模最大基数参考；B0为单轮生成，B1为多轮布尔反馈，B2发送完整反例，B3在192 B或384 B预算内按新增元组覆盖选择见证，并分别设置稳定截断和随机选择对照。")
    add_before(doc,p5,"合成确认集、Bike Sharing和WDBC均包含90个基础实例，每个实例构造C0与C1两个同底表任务，覆盖干净、残留冲突、过度删除、混合错误和歧义五类状态；每个任务独立重复3次，因此每个实验臂在每套数据上产生540条结果，11个实验臂共5940条。Bike Sharing和WDBC分别基于公开数据的显式唯一标识符构造BCNF键依赖，属于受控键冲突修复，不代表任意业务函数依赖。失败状态全部保留在分母中。")
    add_caption(doc,p5,"表1 数据集与实验规模")
    add_table(doc,p5,["数据集","基础实例","C0/C1任务","重复","每臂结果","验证范围"],[
        ["合成确认集","90","180","3","540","单键、多键与复合键"],
        ["Bike Sharing","90","180","3","540","公开表上的受控唯一键冲突"],
        ["WDBC","90","180","3","540","公开表上的受控唯一键冲突"],
    ],[2.6,2.1,2.3,1.6,2.0,5.0])
    add_before(doc,p5,"主指标为形式修复率和有效F1。形式修复成功表示候选满足全部FD且不存在可单独加回的元组；有效F1只对形式检查通过的候选计入事实F1。配对效应先在任务内平均3次重复，再以基础实例为聚类单位进行10000次bootstrap，报告95%置信区间；每个数据集和指标内的预注册比较使用Holm方法校正。BCNF后端实验在Windows 11、Python 3.12.7、20逻辑处理器的单进程环境中计时。")

    add_before(doc,p5,"4.2 主要结果", "Heading 2",False)
    add_before(doc,p5,"表2给出六个代表性实验臂的结果。确定性基线D0在三套数据上均能构造形式合法的极大子集，说明仅以检查通过作为目标不足以证明LLM的必要性。模型的价值应由事实保留质量、反馈纠错效率及业务证据利用能力共同判断。B0的形式修复率仅为7.4%～21.1%；加入多轮布尔反馈后，B1提高到49.1%～59.4%；完整反例B2进一步达到90.2%～93.3%。")
    rows=[]
    for ds in DATASET_LABEL:
        for arm in ["D0","B0","B1","B2","B3_192","B3_384"]:
            r=find_main(ds,arm); c=find_cost(ds,arm)
            rows.append([DATASET_LABEL[ds],ARM_LABEL[arm],pct(r["repair_rate"]),pct(r["valid_f1"]),c["logical_calls"]])
    add_caption(doc,p5,"表2 代表性方法的形式修复率与有效F1（%）")
    add_table(doc,p5,["数据集","方法","修复率","有效F1","逻辑调用数"],rows,[2.7,4.3,2.1,2.1,2.5])
    add_figure(doc,p5,ASSET_DIR/"fig_repair_rate.png","图3 三组数据上的形式修复率 / Fig.3 Formal repair rates on three datasets")
    add_before(doc,p5,"B2相对B1的配对提升在三套数据上均成立：合成确认集提高44.3个百分点（95% CI [36.7, 51.7]），Bike Sharing提高35.7个百分点（95% CI [28.7, 42.8]），WDBC提高32.6个百分点（95% CI [25.6, 39.8]）；三项Holm校正后p值均为0.0032。该结果支持的结论是“可定位错误原因的结构化反例比单一通过/失败信号更能帮助模型完成形式纠错”，不能进一步推出模型选择的记录一定符合客观事实。")
    effect_rows=[]
    for ds in DATASET_LABEL:
        r=find_effect(ds,"B2-B1")
        effect_rows.append([DATASET_LABEL[ds],"B2−B1",f"{float(r['absolute_difference'])*100:+.1f}",f"[{float(r['ci95_low'])*100:+.1f}, {float(r['ci95_high'])*100:+.1f}]",f"{float(r['p_holm']):.4f}"])
    for ds in DATASET_LABEL:
        r=find_effect(ds,"B2:C1-C0")
        effect_rows.append([DATASET_LABEL[ds],"B2: C1−C0",f"{float(r['absolute_difference'])*100:+.1f}",f"[{float(r['ci95_low'])*100:+.1f}, {float(r['ci95_high'])*100:+.1f}]",f"{float(r['p_holm']):.4f}"])
    add_caption(doc,p5,"表3 完整反例与来源证据的配对效应")
    add_table(doc,p5,["数据集","比较","差异/百分点","95% CI","pHolm"],effect_rows,[3.0,3.2,2.6,4.0,2.0])
    add_figure(doc,p5,ASSET_DIR/"fig_paired_effects.png","图4 主要配对效应 / Fig.4 Main paired effects")
    add_before(doc,p5,"事实保留结果没有呈现单一方向。D0在合成确认集上的有效F1为93.7%，高于B2的87.9%；在Bike Sharing和WDBC上，B2的有效F1分别为86.3%和88.0%，略高于D0的85.6%和86.1%。这些点估计未构成预注册的跨方法事实质量显著性结论。来源证据C1仅在合成确认集上使B2修复率提高9.6个百分点（95% CI [4.4, 15.2]，pHolm=0.0032）；Bike Sharing的差异为+2.6个百分点，WDBC为−3.3个百分点，区间均跨0。这说明简单标量来源证据的作用依赖任务构造。")

    add_before(doc,p5,"4.3 反馈预算、失败与费用", "Heading 2",False)
    add_before(doc,p5,"在192 B预算下，覆盖优先策略与稳定截断在三套数据上的修复率完全相同；相对随机选择的差异也很小。在384 B预算下，覆盖优先相对两个对照的差异均未通过Holm校正，置信区间跨0。图5显示各比较围绕零点分布。因此，现有实验不支持“覆盖优先必然优于简单压缩”的主张。更稳妥的解释是，在当前任务规模和见证冗余度下，反馈字节数比见证排序规则更重要。")
    add_figure(doc,p5,ASSET_DIR/"fig_budget_effect.png","图5 覆盖优先策略相对预算对照的修复率差异 / Fig.5 Repair-rate differences against budget controls")
    add_before(doc,p5,"失败记录表明，B1常因重复状态、轮数上限或上下文上限停止；B2提供具体反例后，三套数据的失败数分别从275、246和219降至36、53和43。所有失败均保留在540条结果的分母中。Bike矩阵中一组首轮网络中断通过独立定点补跑恢复，对应9条结果单独封存；派生矩阵的540组首轮响应配对检查全部通过。")
    cost_rows=[]
    actual={"synthetic_confirmation":2.925210,"semireal_bike":3.697670,"semireal_wdbc":3.291187}
    for ds in DATASET_LABEL:
        b1=find_cost(ds,"B1"); b2=find_cost(ds,"B2")
        cost_rows.append([DATASET_LABEL[ds],b1["logical_calls"],b2["logical_calls"],f"{float(b2['mean_elapsed_seconds']):.2f}",f"{actual[ds]:.6f}"])
    add_caption(doc,p5,"表4 模型调用、时延与实际账本费用")
    add_table(doc,p5,["数据集","B1调用","B2调用","B2平均用时/s","矩阵费用/元"],cost_rows,[3.3,2.3,2.3,3.2,3.0])
    add_before(doc,p5,"三组确认与外部验证矩阵的实际账本费用合计9.914067元，其中合成、Bike Sharing和WDBC分别为2.925210、3.697670和3.291187元。账本费用受缓存和共享首轮调用影响，因此表4用于复现实验成本，不把各实验臂的归因等价成本误写为实际支付金额。GLM-4.7-Flash和GLM-4.5-Air的新鲜单请求探针持续返回HTTP 429，本文未形成第二模型完整矩阵，相关结果不能写成跨模型复现。")

    add_before(doc,p5,"4.4 BCNF检查后端", "Heading 2",False)
    add_before(doc,p5,"离线微基准覆盖单键、双键和复合键结构、不同冲突比例及不同记录规模。BCNF共享键索引、通用FD哈希、通用检查后端和独立参考判定在6000个候选状态上的最低一致率为100%，支持实现正确性。图6给出复合键、冲突比例0.05时的规模扩展。")
    add_figure(doc,p5,ASSET_DIR/"fig_bcnf_scaling.png","图6 复合键场景下的BCNF后端规模扩展 / Fig.6 Scaling of BCNF backends in the composite-key setting")
    bench_rows=[]
    alg_names={"bcnf_index":"BCNF共享键索引","fd_hash":"通用FD哈希","incremental":"增量回放","repair_agent_generic":"通用检查后端"}
    for alg in alg_names:
        r=next(x for x in BENCH if x["family"]=="composite" and x["row_count"]=="50000" and x["conflict_ratio"]=="0.05" and x["algorithm"]==alg)
        bench_rows.append([alg_names[alg],f"{ms(r['median_decision_ns']):.1f}",f"{ms(r['p95_decision_ns']):.1f}",f"{float(r['agreement_rate'])*100:.1f}"])
    add_caption(doc,p5,"表5 5万记录复合键场景的后端结果")
    add_table(doc,p5,["后端","中位判定/ms","P95判定/ms","一致率/%"],bench_rows,[5.2,3.1,3.1,2.7])
    add_before(doc,p5,"在5万记录复合键场景中，BCNF共享键索引的中位判定时间为89.8 ms，通用FD哈希为134.7 ms，通用检查后端为150.3 ms，增量回放为115.3 ms。相对通用检查后端，BCNF专用后端减少40.3%的判定时间。该结果说明BCNF结构能够降低当前实现的重复检查开销，但端到端模型交互仍主要受网络与推理时间支配；增量回放也不应被解释为已经降低在线API等待时间。")

    # Replace conclusion.
    conclusion_index=next(i for i,p in enumerate(doc.paragraphs) if p.text.strip()=="5 结论")
    conclusion_body=doc.paragraphs[conclusion_index+1]
    conclusion_body.text=("本文围绕BCNF关系模式下的子集修复检查，将规范化理论、共享键索引和大模型辅助修复结合起来。第2章保留并明确了函数依赖删除封闭性、BCNF超键唯一性和键值冲突局部性在判定中的不同作用；在此基础上，检查器以冲突元组对解释一致性失败，以可加回元组解释极大性失败。三套冻结数据的确认性结果表明，完整反例相对布尔反馈使形式修复率提高32.6～44.3个百分点，且三组配对比较经Holm校正后均显著。BCNF专用后端在6000个状态上保持100%判定一致，并在5万记录复合键场景中较通用后端减少40.3%的中位判定时间。与此同时，覆盖优先的预算化见证选择没有稳定优于截断或随机，来源证据在两套半真实数据上的效应也不显著。本文的结论限于可信FD、元组删除和单一模型版本；后续工作应补充可用的第二模型、真实业务FD与更强的事实质量基线，并研究无需完整枚举的有界见证生成。")
    conclusion_body.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    conclusion_body.paragraph_format.first_line_indent=Pt(21)
    for r in conclusion_body.runs: set_run_font(r,10.5)

    # Add dataset citations after the existing references.
    refs=[
        "[11] UCI Machine Learning Repository. Bike Sharing Dataset. DOI: 10.24432/C5W894.",
        "[12] UCI Machine Learning Repository. Breast Cancer Wisconsin (Diagnostic). DOI: 10.24432/C5DW2B.",
    ]
    for txt in refs:
        p=doc.add_paragraph(txt,style="Reference")
        for r in p.runs: set_run_font(r,9.0)

    # Journal-like spacing, heading protection, table cell margins, and page number update on open.
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading"):
            p.paragraph_format.keep_with_next=True
            p.paragraph_format.keep_together=True
        if p.style.name == "Figure Caption":
            p.paragraph_format.keep_with_next=True
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells: set_cell_margins(cell)
    settings=doc.settings._element
    upd=settings.find(qn("w:updateFields"))
    if upd is None: upd=OxmlElement("w:updateFields"); settings.append(upd)
    upd.set(qn("w:val"),"true")
    doc.core_properties.title="BCNF条件下的子集修复检查与大模型反例反馈方法"
    doc.core_properties.subject="关系数据库规范化理论、子集修复检查与大模型反例反馈"
    doc.core_properties.comments="由冻结实验结果生成的实验增强稿；原论文及理论保留版均未覆盖。"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
